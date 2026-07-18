"""June F2/F3 eligibility manifest (founder §6) + candidate chronological designs (§7).

OUTCOME-BLIND: uses Tennis-Data pre-June rows (dates/names/comments/surfaces — the
winner IDENTITY is never read here; history counts use row PRESENCE under the label
policy, counting each player's appearances regardless of result) and June Betfair
runner NAMES (metadata). No June outcome, no odds value.

Label policy (founder §4, label-policy-v1): eligible history rows = Comment COMPLETED
only. Prior-history band per market = band(min over both players of pre-June completed
appearances). Policies: source vintage-2026-07-18, identity td-norm-v1,
eligibility-policy-v1.
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import date, datetime

sys.path.insert(0, "/home/user/Moneymaker")
import openpyxl
import xlrd
from decimal import Decimal

from l8_evidence.sample_size import PowerAssumptions, derive_sample_size
from sport_tennis.identity_bridge import build_bridge

ROOT = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(ROOT, "raw", "vintage-2026-07-18")
BOUNDARY = date(2026, 6, 1)
BANDS = [(0, 0, "0"), (1, 4, "1-4"), (5, 9, "5-9"), (10, 19, "10-19"), (20, 10**9, "20+")]


def iter_rows(path):
    if path.endswith(".xlsx"):
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = wb.active.iter_rows(values_only=True)
        hdr = [str(c) for c in next(rows)]
        for r in rows:
            yield hdr, r
    else:
        wb = xlrd.open_workbook(path)
        ws = wb.sheet_by_index(0)
        hdr = [str(c.value) for c in ws.row(0)]
        for i in range(1, ws.nrows):
            vals = []
            for c in ws.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    vals.append(datetime(*xlrd.xldate_as_tuple(c.value, wb.datemode)))
                elif c.ctype == xlrd.XL_CELL_EMPTY:
                    vals.append(None)
                else:
                    vals.append(c.value)
            yield hdr, tuple(vals)


def parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for f in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None


def band(n: int) -> str:
    for lo, hi, label in BANDS:
        if lo <= n <= hi:
            return label
    return "20+"


def main() -> None:
    # ---- parse TD completed rows (pre-June), collecting per-player appearance dates ----
    appearances: dict[tuple[str, str], list[date]] = defaultdict(list)  # (tour, name)->dates
    rows_by_block: list[tuple[str, date, str]] = []  # (tour, date, surface) completed only
    names_by_tour: dict[str, set] = {"ATP": set(), "WTA": set()}
    for fn in sorted(os.listdir(RAW)):
        tour = "ATP" if fn.startswith("atp") else "WTA"
        for hdr, row in iter_rows(os.path.join(RAW, fn)):
            cols = {h: (row[i] if i < len(row) else None) for i, h in enumerate(hdr)}
            d = parse_date(cols.get("Date"))
            if d is None or d >= BOUNDARY:
                continue
            cm = str(cols.get("Comment") or "Completed").strip().lower()
            if not cm.startswith("completed"):
                continue  # label-policy-v1: completed only
            w, l = cols.get("Winner"), cols.get("Loser")
            if not w or not l:
                continue
            surf = str(cols.get("Surface") or "").strip()
            rows_by_block.append((tour, d, surf))
            for name in (str(w).strip(), str(l).strip()):
                names_by_tour[tour].add(name)
                appearances[(tour, name)].append(d)

    # ---- bridge to June Betfair names ----
    uni_rows = list(csv.DictReader(open("/home/user/Moneymaker/docs/evidence/pilot-2026-06-tennis/PRIMARY_JUNE_ANALYSIS_UNIVERSE_MANIFEST.csv")))
    uni = {r["market_id"]: ("STRICT" in r["cohort_tags"]) for r in uni_rows if r["universe_membership"] == "PRIMARY_JUNE_SINGLES_UNIVERSE"}
    market_names: dict[str, list[str]] = {}
    with open(os.path.join(ROOT, "..", "pilot-data", "audit", "full_corpus_metadata.jsonl")) as f:
        for line in f:
            r = json.loads(line)
            if r["market_id"] in uni and r.get("runner_names"):
                market_names[r["market_id"]] = [n.strip() for n in r["runner_names"]]
    all_bf = sorted({n for ns in market_names.values() for n in ns})
    res = build_bridge(
        td_names_by_tour={"ATP": sorted(names_by_tour["ATP"]), "WTA": sorted(names_by_tour["WTA"])},
        betfair_full_names=all_bf,
        source_vintage="vintage-2026-07-18",
    )
    mapped = {m.betfair_alias.display_name: m for m in res.mappings}
    unresolved_reason = {u.name: u.reason.value for u in res.unresolved if u.name in set(all_bf)}

    # ---- per-market eligibility manifest ----
    manifest = []
    excl_funnel: Counter = Counter()
    band_counts: Counter = Counter()
    f2 = f2_strict = 0
    for mid in sorted(market_names):
        names = market_names[mid]
        strict = uni[mid]
        mapped_flags = [n in mapped for n in names]
        both = all(mapped_flags)
        if both:
            counts = []
            for n in names:
                m = mapped[n]
                counts.append(len(appearances[(m.source.tour, m.source.raw_name)]))
            prior_band = band(min(counts))
            history = all(c > 0 for c in counts)
            f2_ok = True  # cold-start supported: identity resolved => F2 can emit
            reason = "ELIGIBLE_F2"
            f2 += 1
            f2_strict += strict
            band_counts[prior_band] += 1
        else:
            prior_band = None
            history = False
            f2_ok = False
            missing = [unresolved_reason.get(n, "NO_SOURCE_MATCH") for n, ok in zip(names, mapped_flags) if not ok]
            reason = "EXCLUDED_IDENTITY:" + "|".join(sorted(set(missing)))
        excl_funnel[reason.split(":")[0]] += 1
        manifest.append(
            {
                "market_id": mid,
                "both_mapped": both,
                "pre_june_history_available": history,
                "f2_can_emit": f2_ok,
                "f3_can_emit": False,
                "f3_blocker": "JUNE_SURFACE_SOURCE_MISSING",
                "prior_match_band_min_player": prior_band,
                "cohort": "STRICT" if strict else "PRIMARY_ONLY_TIER",
                "reason": reason,
                "source_vintage": "vintage-2026-07-18",
                "identity_policy": "td-norm-v1",
                "eligibility_policy": "eligibility-policy-v1",
                "label_policy": "label-policy-v1-completed-only",
            }
        )
    body = "\n".join(json.dumps(m, sort_keys=True) for m in manifest) + "\n"
    digest = hashlib.sha256(body.encode()).hexdigest()
    with open(os.path.join(ROOT, "JUNE_F2F3_ELIGIBILITY_MANIFEST.jsonl"), "w") as f:
        f.write(body)

    summary = {
        "markets": len(manifest),
        "both_player_mapped": sum(1 for m in manifest if m["both_mapped"]),
        "f2_eligible": f2,
        "f2_eligible_strict": f2_strict,
        "f2_eligible_primary_only_tier": f2 - f2_strict,
        "f3_eligible": 0,
        "f3_blocker": "JUNE_SURFACE_SOURCE_MISSING (governed June-surface source required)",
        "prior_history_bands_f2_eligible": dict(band_counts),
        "exclusion_funnel": dict(excl_funnel),
        "coverage_ledger_note": "760 source-unmatched players and 54 ambiguity/homonym refusals remain visible in identity_resolution_report.json; nothing dropped",
        "manifest_sha256": digest,
    }
    json.dump(summary, open(os.path.join(ROOT, "JUNE_F2F3_ELIGIBILITY_SUMMARY.json"), "w"), indent=1, sort_keys=True)
    print(json.dumps(summary, indent=1, sort_keys=True))

    # ---- candidate designs ----
    def block_counts(lo: date, hi: date):
        c: Counter = Counter()
        s: Counter = Counter()
        for tour, d, surf in rows_by_block:
            if lo <= d <= hi:
                c[tour] += 1
                s[surf] += 1
        return c, s

    designs = []
    for name, warm_hi, dev_lo, dev_hi, val_lo, val_hi in (
        ("A", date(2020, 12, 31), date(2021, 1, 1), date(2025, 8, 31), date(2025, 9, 1), date(2026, 5, 31)),
        ("B", date(2018, 12, 31), date(2019, 1, 1), date(2025, 5, 31), date(2025, 6, 1), date(2026, 5, 31)),
    ):
        warm, _ = block_counts(date(2000, 1, 1), warm_hi)
        dev, dev_surf = block_counts(dev_lo, dev_hi)
        val, val_surf = block_counts(val_lo, val_hi)
        oof_n = sum(dev.values())
        val_n = sum(val.values())
        eval_n = oof_n + val_n
        power_rows = {}
        for k, n_req in (("alpha_0.05_single", 26600), ("alpha_0.05_over_6", 41039)):
            power_rows[k] = {"n_required": n_req, "n_available_oof_plus_validation": eval_n, "adequate": eval_n >= n_req}
        designs.append(
            {
                "design": name,
                "warm_up": {"through": str(warm_hi), "matches": dict(warm)},
                "development_oof": {"from": str(dev_lo), "to": str(dev_hi), "matches": dict(dev), "surfaces": dict(dev_surf)},
                "validation": {"from": str(val_lo), "to": str(val_hi), "matches": dict(val), "surfaces": dict(val_surf)},
                "oof_prediction_count": oof_n,
                "validation_count": val_n,
                "evaluation_total": eval_n,
                "power_vs_declared": power_rows,
            }
        )
    json.dump({"designs": designs}, open(os.path.join(ROOT, "CANDIDATE_DESIGNS.json"), "w"), indent=1, sort_keys=True)
    for d in designs:
        print(f"Design {d['design']}: warm={sum(d['warm_up']['matches'].values()):,} "
              f"dev/OOF={d['oof_prediction_count']:,} val={d['validation_count']:,} "
              f"eval_total={d['evaluation_total']:,} "
              f"adequate@0.05={d['power_vs_declared']['alpha_0.05_single']['adequate']} "
              f"@0.05/6={d['power_vs_declared']['alpha_0.05_over_6']['adequate']}")


if __name__ == "__main__":
    main()
